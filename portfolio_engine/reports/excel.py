"""Excel workbook output. One workbook per build: a summary matrix, allocation tables,
tactical signals, all holdings in long format, a sheet per portfolio, the universe with
latest prices, and the assumptions the build used."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..builder import Portfolio
from ..config import Profiles, Settings
from ..market_data import MarketData
from ..signals import TacticalView

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12)
WARN_FILL = PatternFill("solid", fgColor="FFF4CE")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(bottom=THIN)

PCT = "0.00%"
PCT1 = "0.0%"
MONEY = "$#,##0"
MONEY2 = "$#,##0.00"


def _write_table(ws, df: pd.DataFrame, start_row: int, formats: dict[str, str] | None = None, col_widths: dict | None = None) -> int:
    formats = formats or {}
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=str(col))
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, val in enumerate(row, start=1):
            if isinstance(val, float) and pd.isna(val):
                val = None
            c = ws.cell(row=i, column=j, value=val)
            fmt = formats.get(df.columns[j - 1])
            if fmt:
                c.number_format = fmt
            c.border = BORDER
    for j, col in enumerate(df.columns, start=1):
        longest = df[col].astype(str).str.len().max() if len(df) else 0
        longest = 0 if pd.isna(longest) else int(longest)
        width = (col_widths or {}).get(col) or min(max(12, len(str(col)) + 2, longest + 2), 48)
        ws.column_dimensions[get_column_letter(j)].width = width
    return start_row + len(df) + 2


def _sheet_name(pf: Portfolio, profiles: Profiles) -> str:
    ab = {"conservative": "Cons", "moderate": "Mod", "balanced": "Bal", "growth": "Gro", "high_growth": "HiGro"}
    ls = {"early_accumulation": "Early", "mid_accumulation": "Mid", "pre_retirement": "PreRet", "retirement": "Ret"}
    tier = {"starter": "Start", "core": "Core", "established": "Estab", "high": "High"}
    return f"{ab.get(pf.profile_used, pf.profile_used[:5])}-{ls.get(pf.life_stage, pf.life_stage[:5])}-{tier.get(pf.tier, pf.tier[:5])}"[:31]


def write_workbook(path: Path, portfolios: list[Portfolio], profiles: Profiles, settings: Settings, md: MarketData,
                   view: TacticalView | None, universe: pd.DataFrame, prices: dict[str, float], origin: dict[str, str],
                   research: dict | None = None, review: list | None = None, quality: dict | None = None,
                   pds: dict | None = None, esg: dict | None = None) -> Path:
    wb = Workbook()
    labels = {c: v["label"] for c, v in profiles.asset_classes.items()}
    classes = list(profiles.asset_classes)

    # ------------------------------------------------------------ Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Model portfolio matrix"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Built {datetime.now():%Y-%m-%d %H:%M}. Prices as of {md.as_of.date()}."
    ws["A3"] = ("DATA IS SYNTHETIC (offline test mode). Do not use these numbers."
                if md.synthetic else f"Price sources: {', '.join(sorted(set(md.source_by_ticker.values())))}. FX AUD per USD {md.fx_aud_per_usd:.4f}.")
    if md.synthetic:
        ws["A3"].fill = WARN_FILL
        ws["A3"].font = Font(bold=True)
    ws["A4"] = ("Tactical tilts ENABLED: class weights are strategic allocation plus bounded rule-based tilts (see Signals)."
                if view and view.enabled else "Tactical tilts disabled: class weights are the strategic allocation.")
    ws["A5"] = ("Historical figures (volatility, trailing return) describe the past and are not forecasts. Weighted MER excludes "
                "platform administration fees, which are shown separately from the HUB24 rate card. Weighted 3, 5 and 10 year returns are the "
                "weighted average of each holding's annualised return (holdings younger than the period use their asset class index ETF; "
                "the proxy share column says how much of the portfolio that covers). Yields are live trailing dividend yields where available.")
    rows = []
    for pf in portfolios:
        m = {k: (None if v is None else v) for k, v in pf.metrics.items()}
        rows.append({
            "Portfolio": pf.id, "Implementation": "Managed portfolio (SMA)" if pf.implementation == "sma" else "Individual holdings", "ESG screen": "on" if (pf.esg or {}).get("screened") else "off", "Risk profile": profiles.risk_profiles[pf.profile_used]["label"],
            "Life stage": profiles.life_stages[pf.life_stage]["label"], "Balance tier": profiles.balance_tiers[pf.tier]["label"],
            "Balance": pf.balance, "Growth": m["growth_pct"] / 100, "Defensive": m["defensive_pct"] / 100,
            "Holdings": m["holdings"], "Weighted MER": m["weighted_mer_pct"] / 100,
            "Yield": (m["weighted_yield_pct"] / 100) if m.get("weighted_yield_pct") is not None else None,
            "Income p.a.": m.get("income_per_year"), "Franking credits p.a.": m.get("franking_credits_per_year"),
            "Grossed-up yield": (m["grossed_up_yield_pct"] / 100) if m.get("grossed_up_yield_pct") is not None else None,
            "Australian equities": (m["aus_equity_pct"] / 100) if m.get("aus_equity_pct") is not None else None,
            "Investment fees p.a.": m["investment_fees_per_year"], "Platform menu": m.get("platform_menu", "choice"),
            "Platform fee p.a.": m["platform_admin_fee_per_year"], "Total ongoing cost": m["total_ongoing_cost_pct"] / 100,
            "Wtd 3y return p.a.": (m.get("weighted_return_3y_pct") or 0) / 100, "Wtd 5y return p.a.": (m.get("weighted_return_5y_pct") or 0) / 100,
            "Wtd 10y return p.a.": (m.get("weighted_return_10y_pct") or 0) / 100, "10y proxy share": (m.get("weighted_return_10y_proxy_share_pct") or 0) / 100,
            "Realised vol (1y)": (m.get("realised_volatility_pct") or 0) / 100,
            "Backtest end value": m.get("backtest_end_value"), "Backtest p.a.": (m.get("backtest_cagr_pct") or 0) / 100,
            "Backtest worst fall": (m.get("backtest_max_drawdown_pct") or 0) / 100, "Backtest stand-in share": (m.get("backtest_stand_in_share_pct") or 0) / 100,
            "Beta ASX 200": m.get("beta_asx200"), "Beta world": m.get("beta_world"), "Corr ASX 200": m.get("correlation_asx200"),
            "Avg pairwise corr": m.get("avg_pairwise_correlation"), "Diversification ratio": m.get("diversification_ratio"),
            "Trailing 1y return": (m.get("trailing_1y_return_pct") or 0) / 100,
            "Warnings": "; ".join(pf.warnings),
        })
    df = pd.DataFrame(rows)
    _write_table(ws, df, 7, {"Balance": MONEY, "Growth": PCT1, "Defensive": PCT1, "Weighted MER": PCT, "Yield": PCT,
                             "Income p.a.": MONEY, "Franking credits p.a.": MONEY, "Grossed-up yield": PCT, "Australian equities": PCT1,
                             "Investment fees p.a.": MONEY, "Platform fee p.a.": MONEY,
                             "Total ongoing cost": PCT, "Realised vol (1y)": PCT1, "Trailing 1y return": PCT1,
                             "Backtest end value": MONEY, "Backtest p.a.": PCT1, "Backtest worst fall": PCT1, "Backtest stand-in share": PCT1,
                             "Wtd 3y return p.a.": PCT1, "Wtd 5y return p.a.": PCT1, "Wtd 10y return p.a.": PCT1, "10y proxy share": PCT1},
                 {"Portfolio": 40, "Warnings": 60, "Life stage": 30, "Balance tier": 32})
    ws.freeze_panes = "B8"

    # ------------------------------------------------------------ Allocation
    ws = wb.create_sheet("Allocation")
    ws["A1"] = "Strategic allocation by risk profile (percent)"
    ws["A1"].font = SECTION_FONT
    saa_rows = [{"Risk profile": rp["label"], **{labels[c]: rp["saa"][c] / 100 for c in classes},
                 "Tolerance band (pp)": rp["tolerance_pp"], "Minimum horizon (years)": rp["min_horizon_years"]}
                for rp in profiles.risk_profiles.values()]
    r = _write_table(ws, pd.DataFrame(saa_rows), 3, {labels[c]: PCT1 for c in classes})
    ws.cell(row=r, column=1, value="Current tactical tilts by risk profile (percentage points, zero-sum)").font = SECTION_FONT
    tilt_rows = []
    for name, rp in profiles.risk_profiles.items():
        saa = {c: float(v) for c, v in rp["saa"].items()}
        t = view.tilts_for(saa, profiles) if view else {c: 0.0 for c in saa}
        tilt_rows.append({"Risk profile": rp["label"], **{labels[c]: t[c] for c in classes}})
    r = _write_table(ws, pd.DataFrame(tilt_rows), r + 2, {labels[c]: "+0.00;-0.00;0" for c in classes})
    ws.cell(row=r, column=1, value="Final asset class weights per portfolio (after tilts, life stage rules and rounding)").font = SECTION_FONT
    fin = [{"Portfolio": pf.id, "Balance": pf.balance, **{labels[c]: pf.class_weights().get(c, 0.0) / 100 for c in classes}}
           for pf in portfolios]
    _write_table(ws, pd.DataFrame(fin), r + 2, {"Balance": MONEY, **{labels[c]: PCT1 for c in classes}}, {"Portfolio": 40})

    # ------------------------------------------------------------ Signals
    ws = wb.create_sheet("Signals")
    ws["A1"] = "Tactical signals"
    ws["A1"].font = SECTION_FONT
    ws["A2"] = ("Trend: price vs 200 day average. Momentum: 12-1 month return less cash. Volatility: 20 day vs 1 year realised "
                "volatility. Score combines them with the weights in profiles.yaml; tilt = score x max tilt, then hysteresis.")
    if view:
        sig_rows = [{"Asset class": labels.get(c, c), "Proxies": ", ".join(s.proxies), "Trend": s.trend, "Momentum": s.momentum,
                     "Volatility ratio": s.volatility, "Score": s.score, "Raw tilt (pp)": s.raw_tilt_pp,
                     "Previous tilt (pp)": s.previous_tilt_pp, "Tilt used (pp)": s.tilt_pp, "Note": s.note}
                    for c, s in view.signals.items()]
        _write_table(ws, pd.DataFrame(sig_rows), 4, {"Trend": PCT, "Momentum": PCT, "Volatility ratio": PCT, "Score": "0.00",
                                                     "Raw tilt (pp)": "+0.00;-0.00;0", "Previous tilt (pp)": "+0.00;-0.00;0",
                                                     "Tilt used (pp)": "+0.00;-0.00;0"}, {"Note": 50})
        cfg = profiles.tactical
        ws.cell(row=6 + len(sig_rows) + 1, column=1, value=f"Max tilt per class {cfg['max_tilt_pp']}pp; max net growth shift "
                f"{cfg.get('max_growth_shift_pp', cfg['max_tilt_pp'])}pp; hysteresis {cfg['hysteresis_pp']}pp; weights {cfg['signal_weights']}.")

    # ------------------------------------------------------------ Holdings (long)
    ws = wb.create_sheet("Holdings")
    hrows = []
    for pf in portfolios:
        for ln in pf.lines:
            rr = (research or {}).get(ln.ticker)
            hrows.append({"Portfolio": pf.id, "Ticker": ln.ticker, "Holding": ln.name, "Asset class": labels.get(ln.asset_class, "Managed portfolio (diversified)"),
                          "Vehicle": ln.vehicle, "Currency": ln.currency, "Weight": ln.weight_pct / 100, "Dollars": ln.dollars,
                          "Units": ln.units, "Price (AUD)": ln.price_aud, "Priced from": ln.priced_from,
                          "MER": ln.mer_pct / 100, "Yield": ln.yield_pct / 100, "Yield source": ln.yield_source,
                          "1y return": (rr.return_1y_pct / 100) if rr and rr.return_1y_pct is not None else None,
                          "3y p.a.": (rr.return_3y_pct_pa / 100) if rr and rr.return_3y_pct_pa is not None else None,
                          "5y p.a.": (rr.return_5y_pct_pa / 100) if rr and rr.return_5y_pct_pa is not None else None,
                          "10y p.a.": (rr.return_10y_pct_pa / 100) if rr and rr.return_10y_pct_pa is not None else None,
                          "Return proxies": ", ".join(f"{k}: {v}" for k, v in rr.return_proxy.items()) if rr and rr.return_proxy else "",
                          "Beta ASX 200 (1y)": (pf.metrics.get("holding_beta_asx200") or {}).get(ln.ticker),
                          "Consensus": ln.consensus_label, "Consensus weighting": ln.consensus_multiplier, "Source": ln.source,
                          "ESG review": ((esg or {}).get("review", {}).get(ln.ticker) or {}).get("band", ""),
                          "Documents": ((pds or {}).get(ln.ticker) or {}).get("url", "")})
    labels["__sma__"] = "Managed portfolio (diversified)"
    _write_table(ws, pd.DataFrame(hrows), 1, {"Weight": PCT, "Dollars": MONEY, "Units": "#,##0.##", "Price (AUD)": MONEY2,
                                              "MER": PCT, "Yield": PCT, "1y return": PCT1, "3y p.a.": PCT1, "5y p.a.": PCT1, "10y p.a.": PCT1},
                 {"Portfolio": 40, "Holding": 44})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ------------------------------------------------------------ Per portfolio sheets
    for pf in portfolios:
        ws = wb.create_sheet(_sheet_name(pf, profiles))
        ws["A1"] = f"{profiles.risk_profiles[pf.profile_used]['label']} | {profiles.life_stages[pf.life_stage]['label']} | {profiles.balance_tiers[pf.tier]['label']}"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"Balance ${pf.balance:,.0f}. Prices as of {pf.as_of}." + ("  SYNTHETIC DATA." if pf.synthetic else "")
        if pf.profile_requested != pf.profile_used:
            ws["A3"] = f"Requested {pf.profile_requested}; capped to {pf.profile_used} by life stage rule."
            ws["A3"].fill = WARN_FILL
        m = {k: (float("nan") if v is None else v) for k, v in pf.metrics.items()}
        metric_rows = ([("Managed portfolio", f"{pf.sma['chosen']['code']} {pf.sma['chosen']['name']} ({pf.sma['chosen']['manager']})"),
                        ("Shortlist", "; ".join(f"{c['code']} {c['name']} {c['total_fee'] * 100:.2f}%" for c in pf.sma["shortlist"]))]
                       if pf.implementation == "sma" else []) + [
            ("Growth / defensive", f"{m['growth_pct']:.1f}% / {m['defensive_pct']:.1f}%"),
            ("Holdings", m["holdings"]),
            ("Weighted MER", f"{m['weighted_mer_pct']:.2f}%"),
            ("Weighted yield", f"{m.get('weighted_yield_pct', float('nan')):.2f}%"),
            ("Income per year", f"${m.get('income_per_year', float('nan')):,.0f}"),
            ("Investment fees per year", f"${m['investment_fees_per_year']:,.0f}"),
            ("Platform menu", str(m.get("platform_menu", "choice")).capitalize()),
            ("Platform administration fee per year", f"${m['platform_admin_fee_per_year']:,.0f}"),
            ("Total ongoing cost", f"{m['total_ongoing_cost_pct']:.2f}%"),
            ("Initial brokerage estimate", f"${m['initial_brokerage']:,.0f}"),
            ("Realised volatility (covariance, 1y)", f"{m.get('realised_volatility_pct', float('nan')):.1f}%" + ("  (proxy: profile SAA in ETFs)" if pf.implementation == "sma" else "")),
            ("Weighted average holding volatility", f"{m.get('weighted_avg_holding_vol_pct', float('nan')):.1f}%  (overstates portfolio risk; shown for comparison)"),
            ("Trailing 1 year return (historical)", f"{m.get('trailing_1y_return_pct', float('nan')):.1f}%"),
            ("Beta to ASX 200 / world shares (1y)", f"{m.get('beta_asx200', float('nan')):.2f} / {m.get('beta_world', float('nan')):.2f}"),
            ("Correlation to ASX 200; average between holdings", f"{m.get('correlation_asx200', float('nan')):.2f}; {m.get('avg_pairwise_correlation', float('nan')):.2f}"),
            ("Diversification ratio", f"{m.get('diversification_ratio', float('nan')):.2f}"),
            ("ESG screen", ("on: " + "; ".join(f"{c['name']} {c['action']}" + (f" for {c['replacement']}" if c.get("replacement") else "") for c in pf.esg.get("changes", []))) if (pf.esg or {}).get("screened") else "off"),
            (f"Backtest: {pf.balance:,.0f} invested {m.get('backtest_years', 10)} years ago, rebalanced monthly", f"${m.get('backtest_end_value', 0):,.0f} today ({m.get('backtest_cagr_pct', float('nan')):.1f}% p.a.; worst fall {m.get('backtest_max_drawdown_pct', float('nan')):.0f}%; {m.get('backtest_stand_in_share_pct', 0):.0f}% on index stand-ins)"),
            ("Weighted 3y / 5y / 10y return p.a. (average of holdings)", f"{m.get('weighted_return_3y_pct', float('nan')):.1f}% / {m.get('weighted_return_5y_pct', float('nan')):.1f}% / {m.get('weighted_return_10y_pct', float('nan')):.1f}%"
             + (f"  ({m.get('weighted_return_10y_proxy_share_pct', 0):.0f}% of the 10y figure from index proxies)" if m.get('weighted_return_10y_proxy_share_pct') else "")),
        ]
        for i, (k, v) in enumerate(metric_rows, start=5):
            ws.cell(row=i, column=1, value=k).font = Font(bold=True)
            ws.cell(row=i, column=2, value=v)
        r = 5 + len(metric_rows) + 1
        ws.cell(row=r, column=1, value="Asset allocation").font = SECTION_FONT
        cw = pf.class_weights()
        alloc = [{"Asset class": labels[c], "Strategic": pf.saa[c] / 100, "Tilt (pp)": pf.tilts[c],
                  "Target": pf.target_class_weights[c] / 100, "Actual": cw.get(c, 0.0) / 100,
                  "Dollars": cw.get(c, 0.0) / 100 * pf.balance} for c in classes]
        r = _write_table(ws, pd.DataFrame(alloc), r + 1, {"Strategic": PCT1, "Tilt (pp)": "+0.00;-0.00;0", "Target": PCT1, "Actual": PCT1, "Dollars": MONEY})
        ws.cell(row=r, column=1, value="Holdings").font = SECTION_FONT
        def _ret(t, key, period):
            rr = (research or {}).get(t)
            if not rr or getattr(rr, key) is None:
                return None
            return getattr(rr, key) / 100
        def _px(t, period):
            rr = (research or {}).get(t)
            return "index proxy" if rr and rr.return_proxy.get(period) else ""
        hold = [{"Ticker": ln.ticker, "Holding": ln.name, "Asset class": labels.get(ln.asset_class, "Managed portfolio (diversified)"), "Vehicle": ln.vehicle,
                 "Weight": ln.weight_pct / 100, "Dollars": ln.dollars, "Units": ln.units, "Price (AUD)": ln.price_aud,
                 "MER": ln.mer_pct / 100, "Yield": ln.yield_pct / 100, "Income p.a.": ln.dollars * ln.yield_pct / 100,
                 "1y": _ret(ln.ticker, "return_1y_pct", "1y"), "3y p.a.": _ret(ln.ticker, "return_3y_pct_pa", "3y"), "3y note": _px(ln.ticker, "3y"),
                 "5y p.a.": _ret(ln.ticker, "return_5y_pct_pa", "5y"), "5y note": _px(ln.ticker, "5y"),
                 "10y p.a.": _ret(ln.ticker, "return_10y_pct_pa", "10y"), "10y note": _px(ln.ticker, "10y"),
                 "Consensus": ln.consensus_label, "Priced from": ln.priced_from} for ln in pf.lines]
        r = _write_table(ws, pd.DataFrame(hold), r + 1, {"Weight": PCT, "Dollars": MONEY, "Units": "#,##0.##", "Price (AUD)": MONEY2,
                                                         "MER": PCT, "Yield": PCT, "Income p.a.": MONEY, "1y": PCT1, "3y p.a.": PCT1,
                                                         "5y p.a.": PCT1, "10y p.a.": PCT1}, {"Holding": 44})
        if pf.warnings:
            ws.cell(row=r, column=1, value="Notes").font = SECTION_FONT
            for i, w in enumerate(pf.warnings, start=r + 1):
                ws.cell(row=i, column=1, value=w).fill = WARN_FILL
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 44

    # ------------------------------------------------------------ Research
    if research:
        ws = wb.create_sheet("Research")
        ws["A1"] = "Holding research: Yahoo Finance fundamentals and analyst consensus; returns from dividend-adjusted prices"
        ws["A1"].font = SECTION_FONT
        ws["A2"] = ("Consensus mean runs from 1 (Strong Buy) to 5 (Sell) across covering brokers. It scales a holding's weight within "
                    "its asset class by at most the configured amount and flags names for review; it never trades on its own.")
        rr = [{"Ticker": r.ticker, "Name": r.name, "Sector": r.sector, "Industry": r.industry, "Consensus": r.consensus_label,
               "Consensus mean": r.consensus_mean, "Analysts": r.analysts, "Target (native ccy)": r.target_mean,
               "Target upside": (r.target_upside_pct / 100) if r.target_upside_pct is not None else None,
               "1y return": (r.return_1y_pct / 100) if r.return_1y_pct is not None else None,
               "3y return p.a.": (r.return_3y_pct_pa / 100) if r.return_3y_pct_pa is not None else None,
               "5y return p.a.": (r.return_5y_pct_pa / 100) if r.return_5y_pct_pa is not None else None,
               "10y return p.a.": (r.return_10y_pct_pa / 100) if r.return_10y_pct_pa is not None else None,
               "Return proxies": ", ".join(f"{k}: {v}" for k, v in r.return_proxy.items()), "History (years)": r.history_years,
               "Volatility 1y": (r.volatility_1y_pct / 100) if r.volatility_1y_pct is not None else None,
               "Max drawdown 1y": (r.max_drawdown_1y_pct / 100) if r.max_drawdown_1y_pct is not None else None,
               "Dividend yield": (r.dividend_yield_pct / 100) if r.dividend_yield_pct is not None else None,
               "Trailing PE": r.pe_trailing, "Forward PE": r.pe_forward, "Beta": r.beta, "Market cap": r.market_cap,
               "52w low": r.week52_low, "52w high": r.week52_high, "Fetched": r.fetched, "Source": r.source,
               "Description": (r.summary[:900] if r.summary else "")} for r in research.values()]
        _write_table(ws, pd.DataFrame(rr), 4, {"Target upside": PCT1, "1y return": PCT1, "3y return p.a.": PCT1, "5y return p.a.": PCT1, "10y return p.a.": PCT1,
                                               "Volatility 1y": PCT1, "Max drawdown 1y": PCT1, "Dividend yield": PCT,
                                               "Market cap": "#,##0", "Consensus mean": "0.0", "Trailing PE": "0.0", "Forward PE": "0.0", "Beta": "0.00"},
                     {"Name": 40, "Description": 80})
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(rr[0]))}{4 + len(rr)}"

    # ------------------------------------------------------------ Quality review
    if quality:
        ws = wb.create_sheet("Quality")
        ws["A1"] = "Holding quality review: written verdicts (core, satellite, speculative, not recommended)"
        ws["A1"].font = SECTION_FONT
        ws["A2"] = "Opinion, based on the research data and what each business is. Not advice. Re-read when the research changes."
        qrows = [{"Ticker": t, "Name": (research or {}).get(t).name if (research or {}).get(t) else t, "Verdict": v["verdict"], "Note": v["note"], "Reviewed": v["reviewed"]}
                 for t, v in quality.items()]
        _write_table(ws, pd.DataFrame(qrows), 4, None, {"Name": 40, "Note": 110})
        ws.freeze_panes = "A5"

    # ------------------------------------------------------------ Review
    if review:
        ws = wb.create_sheet("Review")
        ws["A1"] = "Holding review: strikes against active holdings, merits for watchlist names"
        ws["A1"].font = SECTION_FONT
        rcfg = settings.raw.get("review", {})
        ws["A2"] = (f"Removal needs {rcfg.get('strikes_to_remove', 3)} strikes; addition needs {rcfg.get('add_requirements', 3)} merits and no strikes. "
                    f"Proposals only unless review.auto_apply is true; `python -m portfolio_engine review --apply` makes up to "
                    f"{rcfg.get('max_changes_per_run', 2)} changes with a {rcfg.get('cooling_off_days', 90)} day cooling-off.")
        rrows = [{"Action": it.action, "Ticker": it.ticker, "Name": it.name, "Asset class": labels.get(it.asset_class, it.asset_class),
                  "Status": it.status, "Score": it.score, "Strikes": "; ".join(it.reasons), "Merits": "; ".join(it.positives),
                  "Blocked until": it.blocked_until} for it in review]
        _write_table(ws, pd.DataFrame(rrows), 4, None, {"Name": 40, "Strikes": 70, "Merits": 70})
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:I{4 + len(rrows)}"

    # ------------------------------------------------------------ SMA menu
    sma_path = settings.root / "config" / "sma_menu.csv"
    if sma_path.exists():
        ws = wb.create_sheet("SMA menu")
        menu = pd.read_csv(sma_path)
        ws["A1"] = f"HUB24 managed portfolio menu (from the Platform, Invest and Adviser Fee Calculator, as of {menu['as_of'].iloc[0] if len(menu) else ''})"
        ws["A1"].font = SECTION_FONT
        cols = ["code", "name", "category", "manager", "benchmark", "inception", "ret_1y", "ret_3y", "ret_5y", "ret_10y", "mgmt_fee", "underlying_fees", "transaction_costs", "total_fee"]
        _write_table(ws, menu[cols], 3, {c: PCT for c in ["ret_1y", "ret_3y", "ret_5y", "ret_10y", "mgmt_fee", "underlying_fees", "transaction_costs", "total_fee"]}, {"name": 48, "manager": 40, "benchmark": 40})
        ws.freeze_panes = "B4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{3 + len(menu)}"

    # ------------------------------------------------------------ Universe and prices
    ws = wb.create_sheet("Universe")
    u = universe.drop(columns=["min_tier_order"], errors="ignore").copy()
    u["latest_price_aud"] = u["ticker"].map(prices)
    u["priced_from"] = u["ticker"].map(origin)
    _write_table(ws, u, 1, {"latest_price_aud": MONEY2}, {"name": 44, "notes": 40})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ------------------------------------------------------------ Assumptions
    ws = wb.create_sheet("Assumptions")
    ws["A1"] = "Life stage rules"
    ws["A1"].font = SECTION_FONT
    ls_rows = [{"Life stage": v["label"], "Maximum profile": v["max_profile"], "Default profile": v["default_profile"],
                "Cash floor (pp)": v["cash_floor_pp"], "Income preference": v["income_preference"],
                "Franking preference": v["franking_preference"], "Home bias (pp)": v["home_bias_pp"]} for v in profiles.life_stages.values()]
    r = _write_table(ws, pd.DataFrame(ls_rows), 2)
    ws.cell(row=r, column=1, value="Balance tier rules").font = SECTION_FONT
    t_rows = [{"Tier": v["label"], "Minimum balance": v["min_balance"], "Representative balance": v["representative_balance"],
               "Allowed vehicles": ", ".join(v["allowed_vehicles"]), "Max holdings": v["max_holdings"],
               "Max per class": v["max_per_class"], "Min holding $": v["min_holding_dollars"], "Brokerage $": v["brokerage_dollars"]}
              for v in profiles.balance_tiers.values()]
    r = _write_table(ws, pd.DataFrame(t_rows), r + 1, {"Minimum balance": MONEY, "Representative balance": MONEY, "Min holding $": MONEY})
    plat = settings.raw.get("platform", {})
    menus = plat.get("menus") or {"choice": plat}
    ws.cell(row=r, column=1, value=f"Platform fees: {plat.get('name', 'n/a')} rate card {plat.get('rate_card_date', '')}").font = SECTION_FONT
    r += 1
    ws.cell(row=r, column=1, value=f"Menu rules: listed shares or ETFs use the Choice menu; managed-portfolio-only accounts use Core, or Discover under ${plat.get('discover_max_balance', 0):,.0f}.")
    r += 1
    for key, mc in menus.items():
        ws.cell(row=r, column=1, value=mc.get("label", key.capitalize())).font = SECTION_FONT
        p_rows = [{"Balance up to": b["up_to"] if b["up_to"] is not None else "above", "Rate": b["rate"]} for b in mc.get("bands", [])]
        r = _write_table(ws, pd.DataFrame(p_rows), r + 1, {"Rate": PCT})
        ws.cell(row=r, column=1, value=f"Account keeping fee ${mc.get('account_keeping_fee', 0)}; admin fee min ${mc.get('min_admin_fee', 0)} "
                f"max ${mc.get('max_admin_fee', 0)}; expense recovery {mc.get('expense_recovery_rate', 0):.3%} capped at ${mc.get('expense_recovery_cap', 0)}."
                + (f" {mc['note']}" if mc.get("note") else ""))
        r += 2
    for src in plat.get("sources", []):
        ws.cell(row=r, column=1, value=f"Source: {src}")
        r += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
