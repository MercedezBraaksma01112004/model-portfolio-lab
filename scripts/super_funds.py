"""Super fund comparison data for the website's Compare page, from APRA's Comprehensive Product Performance Package.

APRA publishes, each August, every MySuper product and every choice investment option in a trustee-directed
product with its fees and costs at five balances, net returns over 3, 5, 7 and 10 years, the performance test
result, the strategic growth allocation, assets and member numbers. It is official, covers every fund, and is
licensed under Creative Commons Attribution 3.0 Australia, so it can be republished with attribution.

    python scripts/super_funds.py            # refresh if the saved data is more than a week old
    python scripts/super_funds.py --force    # refresh now

Writes data/super_funds.json (kept in the repository, so a failed download never empties the page).
"""
from __future__ import annotations

import io
import json
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "super_funds.json"
PAGE = "https://www.apra.gov.au/superannuation-product-performance"
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0 Safari/537.36"
BALANCES = [10000, 25000, 50000, 100000, 250000]

SHEETS = {   # sheet name -> kind shown on the page
    "MySuper Products": "MySuper",
    "Non-Platform TDPs": "Choice (diversified)",
    "Platform TDPs": "Platform (diversified)",
    "Non-Platform EDPs": "Choice (single sector)",
}


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def find_links() -> dict:
    html = requests.get(PAGE, headers={"User-Agent": UA}, timeout=40).text
    links = re.findall(r'href="([^"]+CPPP[^"]+\.xlsx)"', html)
    out = {}
    for l in links:
        url = l if l.startswith("http") else "https://www.apra.gov.au" + l
        if "mysuper" in l.lower():
            out.setdefault("mysuper", url)
        elif "choice" in l.lower():
            out.setdefault("choice", url)
    if not out:
        raise RuntimeError("no CPPP workbooks linked from " + PAGE)
    return out


def header_index(rows: list[tuple]) -> int:
    for i, r in enumerate(rows[:12]):
        cells = [str(c or "") for c in r]
        if "RSE name" in cells and any("product name" in c.lower() for c in cells):
            return i
    raise RuntimeError("header row not found")


def col(headers: list[str], *starts: str) -> int | None:
    for s in starts:
        for j, h in enumerate(headers):
            if h.lower().startswith(s.lower()):
                return j
    return None


def parse_sheet(ws, kind: str, as_of_hint: str) -> tuple[list[dict], str]:
    rows = list(ws.iter_rows(values_only=True))
    hi = header_index(rows)
    headers = [re.sub(r"\s+", " ", str(c or "")).strip() for c in rows[hi]]
    as_of = as_of_hint
    for r in rows[:hi]:
        for c in r:
            m = re.search(r"as at\s+(\d{1,2} \w+ \d{4})", str(c or ""), re.I)
            if m:
                as_of = m.group(1)
    c = {
        "fund": col(headers, "RSE name"), "licensee": col(headers, "RSE licensee"), "offer": col(headers, "Public offer status"),
        "product": col(headers, "MySuper product name", "Superannuation product name"), "menu": col(headers, "Investment menu name"),
        "option": col(headers, "Investment option name"), "stage": col(headers, "Lifecycle stage name"), "open": col(headers, "Open or Closed", "Open or closed"),
        "assets": col(headers, "Member assets"), "accounts": col(headers, "Member accounts"), "growth": col(headers, "Strategic growth asset allocation"),
        "growth_cat": col(headers, "Strategic growth asset allocation category"), "test": col(headers, "Pass/Fail indicator"),
        "r10": col(headers, "10 year Net Investment Return", "10 year Gross Investment Return"), "r7": col(headers, "7 year Net Investment Return", "7 year Gross Investment Return"),
        "r5": col(headers, "5 year Net Investment Return", "5 year Gross Investment Return"), "r3": col(headers, "3 year Net Investment Return", "3 year Gross Investment Return"),
        "rm10": col(headers, "10 year Net Return ($50,000"), "lifecycle": col(headers, "Single strategy / Lifecycle"),
        "type": col(headers, "MySuper product type"),
    }
    # "Strategic growth asset allocation" also prefixes the category column; take the first that is not the category.
    if c["growth"] == c["growth_cat"]:
        c["growth"] = next((j for j, h in enumerate(headers) if h.lower() == "strategic growth asset allocation"), None)
    admin = [col(headers, f"Administration fees and costs charged (${b:,}") for b in BALANCES]
    total = [col(headers, f"Total fees and costs charged (${b:,}") for b in BALANCES]
    out = []
    last_parent = None
    for r in rows[hi + 1:]:
        g = lambda k: (r[c[k]] if c[k] is not None and c[k] < len(r) else None)
        fund = g("fund")
        if not fund or str(fund).strip() in ("", "x"):
            continue
        product = str(g("product") or "").strip()
        option = str(g("option") or "").strip() if c["option"] is not None else ""
        stage = str(g("stage") or "").strip() if c["stage"] is not None else ""
        rec = {
            "f": str(fund).strip(), "p": product, "m": str(g("menu") or "").strip() if c["menu"] is not None else "",
            "o": option or (stage and f"Lifecycle: {stage}") or ("MySuper" if kind == "MySuper" else ""),
            "k": kind, "st": stage, "lc": (str(g("lifecycle") or "").lower().startswith("lifecycle")) if c["lifecycle"] is not None else False,
            "open": (str(g("open") or "Open").strip().lower() != "closed"), "po": str(g("offer") or "").strip().lower().startswith("public"),
            "a": (num(g("assets")) * 1000) if num(g("assets")) is not None else None, "n": num(g("accounts")),
            "g": num(g("growth")), "gc": str(g("growth_cat") or "").strip(),
            "pt": str(g("test") or "").strip().replace("*", " (short history)").strip(),
            "r3": num(g("r3")), "r5": num(g("r5")), "r7": num(g("r7")), "r10": num(g("r10")), "rm10": num(g("rm10")),
            "fa": [num(r[j]) if j is not None and j < len(r) else None for j in admin],
            "ft": [num(r[j]) if j is not None and j < len(r) else None for j in total],
        }
        if kind == "Platform (diversified)":
            rec["basis"] = "gross of administration fees"   # platform returns are gross investment returns net of investment fees
        # Lifecycle MySuper: the parent row carries the product's test result and size; stages carry returns and fees.
        if kind == "MySuper" and rec["lc"] and not stage:
            last_parent = rec
            continue
        if kind == "MySuper" and stage and last_parent and last_parent["p"] == product and last_parent["f"] == rec["f"]:
            rec["pt"] = rec["pt"] or last_parent["pt"]
            rec["a"] = rec["a"] or None
            rec["parent_a"] = last_parent["a"]
        if all(v is None for v in (rec["r3"], rec["r5"], rec["r7"], rec["r10"])) and all(v is None for v in rec["ft"]):
            continue
        out.append(rec)
    return out, as_of


def build() -> dict:
    links = find_links()
    rows, as_of = [], ""
    for key, url in links.items():
        print(f"downloading {url}")
        data = requests.get(url, headers={"User-Agent": UA}, timeout=180).content
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet, kind in SHEETS.items():
            if sheet in wb.sheetnames:
                got, as_of = parse_sheet(wb[sheet], kind, as_of)
                print(f"  {sheet}: {len(got)} rows")
                rows += got
    if len(rows) < 200:
        raise RuntimeError(f"only {len(rows)} rows parsed; the workbook layout may have changed")
    return {"generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "as_of": as_of, "balances": BALANCES,
            "source": "APRA Comprehensive Product Performance Package", "source_page": PAGE, "files": links,
            "licence": "Creative Commons Attribution 3.0 Australia (APRA)", "rows": rows}


def main(argv: list[str]) -> int:
    if OUT.exists() and "--force" not in argv:
        try:
            old = json.loads(OUT.read_text())
            age = time.time() - datetime.fromisoformat(old["generated_at"]).timestamp()
            if age < 7 * 86400:
                print(f"super fund data is {age / 86400:.1f} days old; keeping it")
                return 0
        except Exception:  # noqa: BLE001
            pass
    try:
        data = build()
    except Exception as e:  # noqa: BLE001
        print(f"super fund data not refreshed: {e}", file=sys.stderr)
        return 0 if OUT.exists() else 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, separators=(",", ":")))
    print(f"wrote {len(data['rows'])} products and options as at {data['as_of']} to {OUT} ({OUT.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
